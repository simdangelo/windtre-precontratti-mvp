"""Genera il template Sintesi Fisse (per i test MVP) da un documento reale.

Il template contiene solo:
- intestazione WINDTRE + bullet introduttivi;
- tabella intro (statica, valori a merge field);
- "Servizio/i e apparecchiature" (statica, valori a merge field; il bollino
  della tecnologia è inline dopo il valore di "Servizi");
- "Velocità del servizio Internet e mezzi di ricorso" (statica, valori a merge field);
- "Prezzo" (dinamica, 2 colonne);
- "Costi" (dinamica, 1 colonna);
- "Durata, rinnovo e risoluzione" (statica, valori a merge field).

Niente note e niente altre sezioni.

Uso: .venv/bin/python scripts/build_sintesi_template.py
"""

from pathlib import Path
from shutil import copy

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "SuperInternetCasa5G-FWAIndoor-20,99-CambioTecnologia-28102025.docx"
DEST = ROOT / "templates" / "sintesi_fisse_template.docx"

# Indici dei figli del body da TENERE (riferiti al documento sorgente):
# intestazione, bullet, tabella intro, servizi, velocità, prezzo, costi, durata.
KEEP_BODY_INDICES = (
    list(range(0, 13))  # Wind Tre... + bullet + tabella intro + spaziatore
    + [14, 15, 16, 17]  # Servizio/i e apparecchiature
    + [19, 20, 21, 22]  # Velocità
    + [26, 27, 28, 30]  # Prezzo (28 = tabella)
    + [31, 32, 33, 34]  # Costi (33 = tabella)
    + [49, 50, 51, 52]  # Durata (51 = tabella)
)

MERGE_FIELDS = {
    (0, 1): "{{ campo.nome_offerta }}",
    (0, 3): "{{ campo.disponibilita }}",
    (1, 1): "{{ campo.servizi }} {{ bollino }}",
    (1, 3): "{{ campo.traffico_internet }}",
    (1, 5): "{{ campo.traffico_fonia }}",
    (1, 7): "{{ campo.apparecchiature }}",
    (2, 1): "{{ campo.velocita_download }}",
    (2, 2): "{{ campo.velocita_upload }}",
    (7, 1): "{{ campo.durata_minima }}",
    (7, 3): "{{ campo.rinnovo }}",
    (7, 5): "{{ campo.risoluzione_anticipata }}",
    (7, 8): "{{ campo.sblocco_apparato }}",
}

PREZZO_LOOP = [
    ["{%tr for r in prezzo.righe %}", ""],
    ["{{r r.etichetta }}", "{{ r.valore }}"],
    ["{%tr endfor %}", ""],
]

COSTI_LOOP = [
    ["{%tr for r in costi.righe %}"],
    ["{{r r.testo }}"],
    ["{%tr endfor %}"],
]


def set_paragraph_text(paragraph, text):
    for child in list(paragraph._element):
        if child.tag != qn("w:pPr"):
            paragraph._element.remove(child)
    paragraph.add_run(text)


def set_cell_text(cell, text):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    set_paragraph_text(cell.paragraphs[0], text)


def rebuild_table(table, rows_cells):
    for tr in table._tbl.findall(qn("w:tr")):
        table._tbl.remove(tr)
    for texts in rows_cells:
        row = table.add_row()
        for i, text in enumerate(texts):
            set_cell_text(row.cells[i], text)


def prune_body(doc, keep_indices):
    body = doc.element.body
    children = list(body.iterchildren())
    assert len(children) == 62, f"figli attesi: 62, trovati: {len(children)}"
    keep = set(keep_indices)
    for i, child in enumerate(children):
        if i in keep:
            continue
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def main():
    DEST.parent.mkdir(parents=True, exist_ok=True)
    copy(SOURCE, DEST)
    doc = Document(DEST)

    tables = doc.tables
    assert len(tables) == 8, f"tabelle attese: 8, trovate: {len(tables)}"

    for (table_idx, row_idx), placeholder in MERGE_FIELDS.items():
        set_cell_text(tables[table_idx].rows[row_idx].cells[0], placeholder)

    rebuild_table(tables[3], PREZZO_LOOP)  # Prezzo, 2 colonne
    rebuild_table(tables[4], COSTI_LOOP)  # Costi, 1 colonna

    prune_body(doc, KEEP_BODY_INDICES)

    doc.save(DEST)
    print(f"Creato: {DEST} (sezioni superflue e note rimosse)")


if __name__ == "__main__":
    main()
