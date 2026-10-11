import uuid
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.models import GeneratedDocument, TemplateBlock, Voce
from tests.conftest import create_offer


def _payload(db_session) -> dict:
    voce_padre = db_session.scalar(select(Voce).where(Voce.label == "Costo mensile"))
    voce_figlio = db_session.scalar(select(Voce).where(Voce.label == "Per i già clienti"))
    voce_costo = db_session.scalar(
        select(Voce)
        .join(TemplateBlock, Voce.block_id == TemplateBlock.id)
        .where(TemplateBlock.code == "costi", Voce.label == "Costo di attivazione")
    )
    assert voce_padre is not None
    assert voce_figlio is not None
    assert voce_costo is not None
    return {
        "fixed": {"disponibilita": "06-10-2025",
                  "servizi": "Fisso fonia e Internet in tecnologia FWA Indoor",
                  "traffico_internet": "Illimitato", "traffico_fonia": "Chiamate illimitate",
                  "apparecchiature": "Modem Wi-Fi 7", "durata_minima": "24 mesi",
                  "rinnovo": "Rinnovo automatico", "risoluzione_anticipata": "Costo previsto",
                  "sblocco_apparato": "Non applicabile"},
        "rows": [
            {"block": "prezzo", "local_id": "p1", "voce_id": str(voce_padre.id), "value": ""},
            {"block": "prezzo", "local_id": "p2", "parent_local_id": "p1",
             "voce_id": str(voce_figlio.id), "value": "20,99 euro/mese"},
            {"block": "costi", "local_id": "c1", "voce_id": str(voce_costo.id), "value": "10 euro"},
        ],
    }


def _table_texts(docx, marker: str) -> list[str]:
    for table in docx.tables:
        texts = [cell.text for row in table.rows for cell in row.cells]
        if any(marker in text for text in texts):
            return texts
    raise AssertionError(f"Nessuna tabella contenente {marker!r}")


def test_generate_sintesi_docx(client: TestClient, db_session, segment_id, technology_id) -> None:
    offer_id = create_offer(client, "Offerta Da Generare", segment_id, technology_id)
    saved = client.put(f"/api/offers/{offer_id}/sintesi", json=_payload(db_session))
    assert saved.status_code == 200

    response = client.post(f"/api/offers/{offer_id}/generate")
    assert response.status_code == 201
    document_id = response.json()["document_id"]

    record = db_session.get(GeneratedDocument, uuid.UUID(document_id))
    assert record is not None
    path = Path(record.file_path)
    assert path.is_file()

    docx = Document(str(path))
    texts = [cell.text for table in docx.tables for row in table.rows for cell in row.cells]
    assert "Offerta Da Generare" in texts
    assert "06-10-2025" in texts
    assert "20,99 euro/mese" in texts
    assert "Per i già clienti" in texts
    assert "300,00 Mbps in Download" in texts

    prezzo_texts = _table_texts(docx, "Per i già clienti")
    assert "Costo mensile" in prezzo_texts
    assert "20,99 euro/mese" in prezzo_texts
    costi_texts = _table_texts(docx, "Costo di attivazione")
    assert "10 euro" in costi_texts
    assert "20,99 euro/mese" not in costi_texts

    download = client.get(f"/documents/{document_id}/download")
    assert download.status_code == 200
    assert "word" in download.headers["content-type"]


def test_generated_docx_contains_technology_badge(
    client: TestClient, db_session, segment_id, technology_id
) -> None:
    offer_id = create_offer(client, "Offerta Con Bollino", segment_id, technology_id)
    saved = client.put(f"/api/offers/{offer_id}/sintesi", json=_payload(db_session))
    assert saved.status_code == 200

    response = client.post(f"/api/offers/{offer_id}/generate")
    assert response.status_code == 201
    record = db_session.get(
        GeneratedDocument, uuid.UUID(response.json()["document_id"])
    )
    assert record is not None
    docx = Document(record.file_path)
    assert len(docx.inline_shapes) == 1
    assert docx.inline_shapes[0].height == Cm(0.53)
    assert docx.inline_shapes[0].width == Cm(0.53)

    def has_drawing(table) -> bool:
        return any(
            cell._element.findall(".//" + qn("w:drawing"))
            for row in table.rows
            for cell in row.cells
        )

    servizi_table = next(
        table for table in docx.tables
        if any("Servizi" in cell.text for row in table.rows for cell in row.cells)
    )
    assert has_drawing(servizi_table)
    velocita_table = next(
        table for table in docx.tables
        if any("Velocità massimo" in cell.text for row in table.rows for cell in row.cells)
    )
    assert not has_drawing(velocita_table)


def test_generated_docx_without_technology_has_no_badge(
    client: TestClient, db_session, segment_id
) -> None:
    response = client.post("/api/offers", json={
        "name": "Offerta Senza Tecnologia", "segment_id": segment_id,
        "technology_id": None, "fixed": {}, "rows": [],
    })
    assert response.status_code == 201
    offer_id = response.json()["id"]
    generated = client.post(f"/api/offers/{offer_id}/generate")
    assert generated.status_code == 201
    record = db_session.get(
        GeneratedDocument, uuid.UUID(generated.json()["document_id"])
    )
    assert record is not None
    docx = Document(record.file_path)
    assert len(docx.inline_shapes) == 0


def test_generate_skips_when_already_current(
    client: TestClient, db_session, segment_id, technology_id
) -> None:
    offer_id = create_offer(client, "Offerta Skip", segment_id, technology_id)
    assert client.put(
        f"/api/offers/{offer_id}/sintesi", json=_payload(db_session)
    ).status_code == 200
    first = client.post(f"/api/offers/{offer_id}/generate")
    assert first.status_code == 201
    assert first.json()["generated"] is True

    second = client.post(f"/api/offers/{offer_id}/generate")
    assert second.status_code == 200
    assert second.json()["generated"] is False
    assert second.json()["document_id"] == first.json()["document_id"]

    db_session.expire_all()
    assert len(list(db_session.scalars(select(GeneratedDocument)).all())) == 1


def test_generate_force_creates_new_version(
    client: TestClient, db_session, segment_id, technology_id
) -> None:
    offer_id = create_offer(client, "Offerta Forza", segment_id, technology_id)
    assert client.put(
        f"/api/offers/{offer_id}/sintesi", json=_payload(db_session)
    ).status_code == 200
    first = client.post(f"/api/offers/{offer_id}/generate")
    assert first.status_code == 201

    forced = client.post(f"/api/offers/{offer_id}/generate?force=true")
    assert forced.status_code == 201
    assert forced.json()["generated"] is True
    assert forced.json()["document_id"] != first.json()["document_id"]

    db_session.expire_all()
    docs = list(db_session.scalars(select(GeneratedDocument)).all())
    assert len(docs) == 2
    assert sum(1 for document in docs if document.is_current) == 1
    current = next(document for document in docs if document.is_current)
    assert str(current.id) == forced.json()["document_id"]


def test_generate_after_modification_creates_new_version(
    client: TestClient, db_session, segment_id, technology_id
) -> None:
    offer_id = create_offer(client, "Offerta Modificata", segment_id, technology_id)
    payload = _payload(db_session)
    assert client.put(f"/api/offers/{offer_id}/sintesi", json=payload).status_code == 200
    first = client.post(f"/api/offers/{offer_id}/generate")
    assert first.status_code == 201

    payload["fixed"]["disponibilita"] = "07-10-2025"
    assert client.put(f"/api/offers/{offer_id}/sintesi", json=payload).status_code == 200
    second = client.post(f"/api/offers/{offer_id}/generate")
    assert second.status_code == 201
    assert second.json()["generated"] is True
    assert second.json()["document_id"] != first.json()["document_id"]

    db_session.expire_all()
    assert len(list(db_session.scalars(select(GeneratedDocument)).all())) == 2


def test_editor_marks_stale_sintesi(
    client: TestClient, db_session, segment_id, technology_id
) -> None:
    offer_id = create_offer(client, "Offerta Stato Sintesi", segment_id, technology_id)
    payload = _payload(db_session)
    assert client.put(f"/api/offers/{offer_id}/sintesi", json=payload).status_code == 200
    assert client.post(f"/api/offers/{offer_id}/generate").status_code == 201

    page = client.get(f"/offers/{offer_id}")
    assert "Aggiornata" in page.text
    assert "Superata" not in page.text
    assert "stale-warning" not in page.text
    assert "window.SINTESI_STALE = false" in page.text

    payload["fixed"]["disponibilita"] = "08-10-2025"
    assert client.put(f"/api/offers/{offer_id}/sintesi", json=payload).status_code == 200
    page = client.get(f"/offers/{offer_id}")
    assert "Superata" in page.text
    assert "stale-warning" in page.text
    assert "window.SINTESI_STALE = true" in page.text

    regenerated = client.post(f"/api/offers/{offer_id}/generate")
    assert regenerated.status_code == 201
    page = client.get(f"/offers/{offer_id}")
    assert "Aggiornata" in page.text
    assert "window.SINTESI_STALE = false" in page.text
    assert "Versioni precedenti (1)" in page.text
